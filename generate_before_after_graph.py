#!/usr/bin/env python3
"""
Before vs After Hop Reduction Visualizer
=========================================
Generates an interactive HTML showing the dependency graph in three modes:
  Before — original graph, all nodes
  Diff   — all nodes visible; nodes-to-be-removed shown in red dashed
  After  — removed nodes hidden, clean surviving graph

Reads:
  output/tidal_dependency_graph.json
  output/MASTER_Hop_Reduction_Recommendations.xlsx  (data-team-validated source of truth,
                                                     "Master Recommendations" sheet)

Writes:
  output/before_after_hop_reduction.html

Usage:
  python generate_before_after_graph.py
"""

import json
import re
from collections import defaultdict, deque
from pathlib import Path

import openpyxl

SCRIPT_DIR  = Path(__file__).resolve().parent
OUTPUT_DIR  = SCRIPT_DIR / "output"
MASTER_XLSX = OUTPUT_DIR / "MASTER_Hop_Reduction_Recommendations.xlsx"
GRAPH_JSON  = OUTPUT_DIR / "tidal_dependency_graph.json"
OUTPUT_HTML = OUTPUT_DIR / "before_after_hop_reduction.html"
D3_PATH     = OUTPUT_DIR / "d3.v7.min.js"

# Canonical "Validation Status" label written by generate_master_recommendations.py
VS_CONSIDER = "\u2705 Consider"


# ── Simulation (tracks which rec removed each job) ────────────────────────────

def simulate_with_tracking(recommendations: list, graph: dict) -> dict:
    """
    Run savings simulation and track which recommendation caused each job removal.
    Returns removed_jobs (sorted list) and job_to_rec (job_id -> rec metadata dict).
    """
    node_map = {n["id"]: n for n in graph["nodes"]}
    root_set  = {nid for nid, n in node_map.items() if n.get("depth", -1) == 0}

    job_to_rec: dict = {}   # job_id -> {rec_id, category, target_table}
    for rec in recommendations:
        if rec.get("verified_hop_savings", 0) > 0:
            jobs = rec.get("affected_jobs", [])
            if len(jobs) > 1:
                candidates = jobs[:-1] if rec.get("category", "").startswith("D.") else jobs[1:]
                for j in candidates:
                    if j not in root_set and j not in job_to_rec:
                        job_to_rec[j] = {
                            "rec_id":       rec.get("id", ""),
                            "category":     rec.get("category", ""),
                            "target_table": rec.get("target_table", ""),
                        }

    removed_set = set(job_to_rec.keys())
    remaining_nodes = {n["id"]: n for n in graph["nodes"] if n["id"] not in removed_set}
    remaining_links  = [
        lnk for lnk in graph["links"]
        if lnk["source"] in remaining_nodes and lnk["target"] in remaining_nodes
    ]

    reverse_map: dict = defaultdict(set)
    for lnk in remaining_links:
        reverse_map[lnk["target"]].add(lnk["source"])

    reachable: set = set(root_set & remaining_nodes.keys())
    queue = deque(reachable)
    while queue:
        nid = queue.popleft()
        for parent in reverse_map.get(nid, set()):
            if parent not in reachable and parent in remaining_nodes:
                reachable.add(parent)
                queue.append(parent)

    connected_depths = [
        remaining_nodes[nid].get("depth", 0)
        for nid in reachable
        if remaining_nodes[nid].get("depth", -1) >= 0
    ]

    return {
        "removed_jobs": sorted(job_to_rec.keys()),
        "job_to_rec":   job_to_rec,
        "original_max_depth":  max((n["depth"] for n in graph["nodes"]), default=0),
        "new_max_depth":       max(connected_depths, default=0),
        "original_node_count": len(graph["nodes"]),
        "new_node_count":      len(remaining_nodes),
    }


# ── Read MASTER recommendations workbook ──────────────────────────────────────

def _safe_int(v) -> int:
    """Parse leading int, tolerating annotated cells like '0 (see overlap note)'."""
    m = re.match(r"\s*(-?\d+)", str(v or ""))
    return int(m.group(1)) if m else 0


def _safe_float(v) -> float:
    m = re.match(r"\s*(-?\d+(?:\.\d+)?)", str(v or ""))
    return float(m.group(1)) if m else 0.0


def read_master_recommendations(path: Path) -> list:
    """
    Read the "Master Recommendations" sheet of MASTER_Hop_Reduction_Recommendations.xlsx
    (data-team-validated, consolidated across all per-RPT recs — the source of truth).
    Returns a flat list of recommendation dicts, each tagged with 'appears_in_rpts' so
    build_enhanced_data() can scope them to the RPT graph(s) they actually apply to.
    """
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb["Master Recommendations"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 3:
        return []

    # Row 0 is a banner, row 1 is the header row, data starts at row 2
    col = {str(h).strip(): i for i, h in enumerate(rows[1]) if h is not None}

    def _get(row, name, default=None):
        idx = col.get(name)
        return row[idx] if idx is not None and idx < len(row) else default

    records = []
    for row in rows[2:]:
        if not row or not _get(row, "Master ID"):
            continue

        category = str(_get(row, "Category") or "").strip()
        raw_jobs = str(_get(row, "Affected Jobs") or "")
        raw_rpts = str(_get(row, "Appears in RPTs") or "")

        records.append({
            "id":                        str(_get(row, "Master ID") or ""),
            "category":                  category,
            "target_table":              str(_get(row, "Target Table") or ""),
            "risk":                      str(_get(row, "Risk") or ""),
            "affected_jobs":             [j.strip() for j in raw_jobs.split("\n") if j.strip()],
            "appears_in_rpts":           {r.strip() for r in raw_rpts.split("\n") if r.strip()},
            "verified_hop_savings":      _safe_int(_get(row, "Verified Hop Savings")),
            "est_runtime_saved_minutes": _safe_float(_get(row, "Est. Time Saved (min)")),
            # F. recs are LLM-only PL/SQL review notes, not graph-verified hop reductions
            "recommendation_type":       "PLSQL_OPTIMISATION_REVIEW" if category.upper().startswith("F.") else "HOP_REDUCTION",
            "validation_status":         str(_get(row, "Validation Status") or "").strip(),
        })

    return records


# ── Build enhanced graph data ─────────────────────────────────────────────────

def build_enhanced_data(all_graphs: dict, master_recs: list) -> dict:
    """
    For each RPT graph, scope the master's '✅ Consider' recommendations via their
    'appears_in_rpts' tag and run the removal simulation.
    Returns enhanced_data[rpt] = {nodes, links, removed_jobs, job_to_rec,
                                   consider_recs, sim_stats}
    """
    consider_recs_all = [r for r in master_recs if r["validation_status"] == VS_CONSIDER]
    enhanced = {}

    for rpt_table, graph in sorted(all_graphs.items()):
        consider_recs = [r for r in consider_recs_all if rpt_table in r["appears_in_rpts"]]

        hop_recs = [r for r in consider_recs if r.get("recommendation_type") == "HOP_REDUCTION"]
        sim      = simulate_with_tracking(hop_recs, graph)

        rt_saved = round(sum(r["est_runtime_saved_minutes"] for r in hop_recs), 2)
        hop_total = sum(r["verified_hop_savings"] for r in hop_recs)

        enhanced[rpt_table] = {
            "nodes":        graph["nodes"],
            "links":        graph["links"],
            "removed_jobs": sim["removed_jobs"],
            "job_to_rec":   sim["job_to_rec"],
            "consider_recs": [
                {
                    "id":           r["id"],
                    "category":     r["category"],
                    "target_table": r["target_table"],
                    "risk":         r["risk"],
                    "affected_jobs": r["affected_jobs"],
                    "verified_hop_savings": r["verified_hop_savings"],
                    "est_runtime_saved_minutes": r["est_runtime_saved_minutes"],
                }
                for r in consider_recs
            ],
            "sim_stats": {
                "original_max_depth":  sim["original_max_depth"],
                "new_max_depth":       sim["new_max_depth"],
                "original_node_count": sim["original_node_count"],
                "new_node_count":      sim["new_node_count"],
                "depth_saved":         sim["original_max_depth"] - sim["new_max_depth"],
                "nodes_saved":         sim["original_node_count"] - sim["new_node_count"],
                "rt_saved_minutes":    rt_saved,
                "consider_rec_count":  len(consider_recs),
                "hop_savings":         hop_total,
            },
        }
        print(f"  {rpt_table}: {len(sim['removed_jobs'])} jobs to remove, "
              f"depth {sim['original_max_depth']}->{sim['new_max_depth']}, "
              f"{rt_saved} min saved")

    return enhanced


# ── HTML generation ───────────────────────────────────────────────────────────

def generate_html(enhanced_data: dict) -> str:
    data_json = json.dumps(enhanced_data, indent=None).replace("</", "<\\/")

    if D3_PATH.exists():
        with open(D3_PATH, "r", encoding="utf-8") as fh:
            d3_inline = fh.read().replace("</", "<\\/")
    else:
        d3_inline = ""
        print("  WARNING: d3.v7.min.js not found — graph will not render")

    parts = []

    parts.append(r"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>Before vs After — Hop Reduction</title>
<style>
* { margin:0; padding:0; box-sizing:border-box; }
body { font-family:'Segoe UI',Tahoma,Geneva,Verdana,sans-serif; background:#0a0a1a; color:#e0e0e0; display:flex; height:100vh; overflow:hidden; }

/* ── Sidebar ── */
#sidebar { width:320px; background:#12122a; border-right:1px solid #2a2a4a; display:flex; flex-direction:column; flex-shrink:0; overflow:hidden; }
#sidebar-header { padding:14px 16px; border-bottom:1px solid #2a2a4a; }
#sidebar-header h2 { font-size:13px; color:#7b8cff; margin-bottom:10px; }
#rpt-select { width:100%; padding:8px 10px; background:#1a1a3a; border:1px solid #3a3a5a; color:#7b8cff; border-radius:6px; font-size:12px; cursor:pointer; font-weight:600; }
#rpt-select:focus { outline:none; border-color:#7b8cff; }
#rpt-select option { background:#1a1a3a; color:#e0e0e0; }

/* ── Summary banner inside sidebar ── */
#sim-summary { padding:10px 16px; background:#0e1a2a; border-bottom:1px solid #2a2a4a; font-size:11px; }
#sim-summary .sim-row { display:flex; gap:12px; flex-wrap:wrap; margin-top:4px; }
.sim-chip { background:#1a2a3a; border:1px solid #3a4a5a; border-radius:12px; padding:3px 10px; font-size:10px; white-space:nowrap; }
.sim-chip.good { border-color:#2a5a3a; color:#69db7c; }
.sim-chip.warn { border-color:#5a3a1a; color:#ffa94d; }
.sim-chip.info { border-color:#2a3a5a; color:#74c0fc; }

/* ── Changes panel ── */
#changes-header { padding:8px 16px; font-size:11px; font-weight:700; color:#7b8cff; text-transform:uppercase; letter-spacing:0.5px; border-bottom:1px solid #2a2a4a; display:flex; justify-content:space-between; align-items:center; cursor:pointer; user-select:none; }
#changes-header:hover { background:#1a1a3a; }
#changes-list { flex:1; overflow-y:auto; padding:4px 0; }
.change-group { padding:4px 0; }
.change-group-hdr { padding:5px 16px; font-size:10px; font-weight:700; color:#aaa; background:#0e0e22; text-transform:uppercase; letter-spacing:0.5px; display:flex; align-items:center; gap:6px; }
.change-item { padding:5px 14px 5px 24px; font-size:11px; cursor:pointer; border-left:3px solid transparent; display:flex; align-items:flex-start; gap:8px; }
.change-item:hover { background:#1a1a3a; }
.change-item.highlighted { background:#1e1e4a; border-left-color:#ff4444; }
.change-item .dot { width:7px; height:7px; border-radius:50%; flex-shrink:0; margin-top:3px; }
.change-item .cname { font-size:11px; font-weight:500; word-break:break-all; }
.change-item .crec { font-size:9px; color:#888; margin-top:1px; }

/* ── Graph area ── */
#graph-container { flex:1; position:relative; overflow:hidden; display:flex; flex-direction:column; }

/* ── Toolbar ── */
#toolbar { padding:8px 14px; background:#0e0e22; border-bottom:1px solid #2a2a4a; display:flex; gap:10px; align-items:center; flex-shrink:0; }
.mode-btn { padding:5px 16px; border-radius:20px; border:1px solid #3a3a5a; background:#1a1a3a; color:#aaa; cursor:pointer; font-size:12px; transition:all 0.15s; }
.mode-btn:hover { background:#2a2a4a; color:#fff; }
.mode-btn.active { background:#7b8cff; border-color:#7b8cff; color:#fff; font-weight:600; }
.mode-btn.active.diff-active { background:#ff5555; border-color:#ff5555; }
.toolbar-sep { width:1px; height:20px; background:#2a2a4a; }
.toolbar-label { font-size:11px; color:#888; }
#toggle-unlinked { accent-color:#7b8cff; cursor:pointer; margin-right:4px; }
#stats-count { font-size:11px; color:#888; margin-left:auto; }

/* ── SVG ── */
svg { flex:1; width:100%; }
.link-path { fill:none; stroke:#5a6aaa; stroke-width:1.8; stroke-opacity:0.45; }
.link-path.highlighted { stroke:#8b9fff; stroke-width:3; stroke-opacity:0.9; }
.link-path.dimmed { stroke-opacity:0.05; }
.link-path.link-removed { stroke:#ff4444; stroke-dasharray:5,3; stroke-opacity:0.3; }
.link-path.link-removed.highlighted { stroke:#ff6666; stroke-opacity:0.8; }
.node-group { cursor:pointer; }
.node-box { rx:6; ry:6; }
.node-box.highlighted { stroke-width:3 !important; stroke:#fff !important; }
.job-label { font-size:9px; fill:#ddd; pointer-events:none; }
.sql-label { font-size:8px; fill:#69db7c; pointer-events:none; font-style:italic; }
.depth-badge { font-size:7px; fill:#fff; pointer-events:none; font-weight:700; }
.remove-badge { font-size:10px; fill:#fff; pointer-events:none; font-weight:700; }
.difw-badge-ba { font-size:7px; fill:#f59e0b; pointer-events:none; font-weight:700; }
.disabled-badge-ba { font-size:7px; fill:#9ca3af; pointer-events:none; font-weight:700; }
.depth-col-label { font-size:11px; fill:#5a5a8a; font-weight:600; text-anchor:middle; }

/* ── Detail panel ── */
#detail-panel { position:absolute; top:12px; right:12px; width:360px; background:#15152e; border:1px solid #2a2a4a; border-radius:8px; padding:16px; display:none; font-size:12px; max-height:80vh; overflow-y:auto; box-shadow:0 4px 20px rgba(0,0,0,0.5); z-index:10; }
#detail-panel h3 { color:#7b8cff; font-size:13px; margin-bottom:10px; word-break:break-all; }
.dp-field { margin-bottom:8px; }
.dp-label { color:#888; font-size:10px; text-transform:uppercase; letter-spacing:0.5px; }
.dp-value { color:#e0e0e0; margin-top:2px; word-break:break-all; font-size:12px; }
.close-btn { position:absolute; top:8px; right:12px; cursor:pointer; color:#888; font-size:18px; }
.close-btn:hover { color:#fff; }
.removed-notice { background:#2a0a0a; border:1px solid #ff4444; border-radius:6px; padding:8px 12px; font-size:11px; color:#ff8888; margin-bottom:10px; }

/* ── Legend ── */
#legend { position:absolute; bottom:12px; right:12px; background:#15152eee; border:1px solid #2a2a4a; border-radius:8px; padding:10px 14px; font-size:11px; z-index:5; }
#legend .leg-item { display:flex; align-items:center; gap:7px; margin:3px 0; }
#legend .leg-dot { width:10px; height:10px; border-radius:2px; }
.leg-removed { border:2px dashed #ff4444 !important; background:rgba(255,68,68,0.15) !important; }

.tooltip { position:absolute; background:#1e1e3e; border:1px solid #3a3a5a; padding:8px 12px; border-radius:6px; font-size:11px; pointer-events:none; display:none; z-index:100; max-width:380px; word-break:break-all; }
#stats-modal { display:none; position:fixed; top:0; left:0; width:100%; height:100%; background:rgba(0,0,0,0.75); z-index:200; justify-content:center; align-items:center; }
#stats-modal.show { display:flex; }
#stats-content { background:#15152e; border:1px solid #2a2a4a; border-radius:12px; padding:24px 32px; max-width:820px; width:92%; max-height:88vh; overflow-y:auto; box-shadow:0 8px 40px rgba(0,0,0,0.6); position:relative; }
#stats-content h2 { color:#7b8cff; font-size:15px; margin-bottom:14px; }
.stats-close { position:absolute; top:12px; right:16px; cursor:pointer; color:#888; font-size:22px; }
.stats-close:hover { color:#fff; }
.stat-cards-ba { display:flex; gap:10px; flex-wrap:wrap; margin-bottom:16px; }
.stat-card-ba { background:#1a1a3a; border:1px solid #2a2a4a; border-radius:8px; padding:12px 14px; flex:1; min-width:110px; text-align:center; }
.stat-card-ba .sc-val { font-size:22px; font-weight:700; color:#7b8cff; }
.stat-card-ba .sc-label { font-size:10px; color:#888; margin-top:3px; text-transform:uppercase; letter-spacing:0.4px; }
.stat-card-ba .sc-delta { font-size:11px; margin-top:4px; font-weight:600; }
.sc-delta.good { color:#69db7c; } .sc-delta.bad { color:#ff8888; } .sc-delta.flat { color:#888; }
.cmp-section-hdr { color:#aaa; font-size:11px; font-weight:700; margin:14px 0 6px; border-bottom:1px solid #2a2a4a; padding-bottom:3px; text-transform:uppercase; letter-spacing:0.5px; }
.cmp-table { width:100%; border-collapse:collapse; font-size:12px; margin-bottom:4px; }
.cmp-table th { text-align:left; color:#888; font-size:10px; text-transform:uppercase; padding:5px 8px; border-bottom:1px solid #2a2a4a; }
.cmp-table td { padding:4px 8px; border-bottom:1px solid #141428; }
.cmp-table td.num { text-align:right; font-variant-numeric:tabular-nums; }
.cmp-table td.good { color:#69db7c; font-weight:600; text-align:right; }
.cmp-table td.cat-lbl { font-weight:500; }
</style>
</head>
<body>

<div id="sidebar">
  <div id="sidebar-header">
    <h2>&#9660; Before vs After Hop Reduction</h2>
    <select id="rpt-select"></select>
  </div>
  <div id="sim-summary">
    <div style="font-size:10px;color:#888;text-transform:uppercase;letter-spacing:0.5px;">Verified savings (Consider recs)</div>
    <div class="sim-row" id="sim-chips"></div>
  </div>
  <div id="changes-header" onclick="toggleChanges()">
    <span id="changes-title">&#9660; Changes (0 jobs)</span>
    <span style="font-size:10px;color:#888;">click to collapse</span>
  </div>
  <div id="changes-list"></div>
</div>

<div id="graph-container">
  <div id="toolbar">
    <span class="toolbar-label">View:</span>
    <button class="mode-btn" id="btn-before" onclick="setMode('before')">Before</button>
    <button class="mode-btn active diff-active" id="btn-diff" onclick="setMode('diff')">&#9432; Diff</button>
    <button class="mode-btn" id="btn-after" onclick="setMode('after')">After</button>
    <div class="toolbar-sep"></div>
    <label style="font-size:11px;color:#ccc;cursor:pointer;">
      <input type="checkbox" id="toggle-unlinked"> Show Unlinked
    </label>
    <span id="stats-count"></span>
    <button class="mode-btn" onclick="showStatsModal()" style="margin-left:auto;">&#9776; Stats</button>
    <button class="mode-btn" id="btn-difw-ba" onclick="cycleDifwModeBa()">&#11041; DIFW Loads</button>
    <button class="mode-btn" id="btn-disabled-ba" onclick="cycleDisabledModeBa()">&#8856; Disabled Jobs</button>
  </div>
  <div id="stats-modal" onclick="if(event.target===this)this.classList.remove('show')">
    <div id="stats-content">
      <span class="stats-close" onclick="document.getElementById('stats-modal').classList.remove('show')">&times;</span>
      <h2 id="stats-modal-title">Before vs After Statistics</h2>
      <div id="stats-modal-body"></div>
    </div>
  </div>
  <svg id="graph-svg"></svg>
  <div id="detail-panel">
    <span class="close-btn" onclick="document.getElementById('detail-panel').style.display='none'">&times;</span>
    <div id="detail-content"></div>
  </div>
  <div id="legend"></div>
  <div class="tooltip" id="tooltip"></div>
</div>

<script>
""")

    parts.append(d3_inline)
    parts.append("\n</" + "script>\n<script>\n")
    parts.append("var ALL_DATA = ")
    parts.append(data_json)
    parts.append(";\n")

    parts.append(r"""
var COLORS = {
  RPT:"#ff6b6b", FCT:"#ffa94d", DIM:"#69db7c", STG:"#74c0fc",
  MV:"#da77f2",  REF:"#ffd43b", MONTH_END:"#868e96",
  SOURCE:"#20c997", OTHER:"#495057"
};
var CAT_LABELS = {
  RPT:"Report Tables", FCT:"Fact Tables", DIM:"Dimension Tables",
  STG:"Staging Tables", MV:"Materialized Views", REF:"Reference Tables",
  MONTH_END:"Month-End Jobs", SOURCE:"Source/Extract", OTHER:"Other"
};
var CAT_ORDER = {RPT:0,FCT:1,MV:2,DIM:3,STG:4,REF:5,SOURCE:6,MONTH_END:7,OTHER:8};

var COL_WIDTH = 320, ROW_HEIGHT = 72, NODE_W = 260, NODE_H = 58, PAD_LEFT = 60, PAD_TOP = 60;

var rptKeys = Object.keys(ALL_DATA).sort();
var mode = 'diff';
var currentRpt = rptKeys[0];
var currentData = null;
var nodeMap = new Map();
var upMap   = new Map();
var downMap = new Map();
var showUnlinked = false;
var changesOpen = true;
var _linkPaths = null, _nodeG = null;
var selectedJobId = null;

var svg  = d3.select("#graph-svg");
var gRoot = svg.append("g");
var zoom = d3.zoom().scaleExtent([0.05,5]).on("zoom", function(e){ gRoot.attr("transform",e.transform); });
svg.call(zoom);

/* ── Populate RPT dropdown ── */
var rptSel = document.getElementById("rpt-select");
rptKeys.forEach(function(rpt){
  var o = document.createElement("option"); o.value = rpt; o.textContent = rpt;
  rptSel.appendChild(o);
});
rptSel.value = currentRpt;
rptSel.addEventListener("change", function(e){ loadRpt(e.target.value); });

/* ── Mode buttons ── */
function setMode(m) {
  mode = m;
  ['before','diff','after'].forEach(function(x){
    var btn = document.getElementById('btn-'+x);
    btn.classList.toggle('active', x===m);
    if (x==='diff') btn.classList.toggle('diff-active', x===m);
  });
  renderGraph();
}

/* ── Toggle changes panel ── */
function toggleChanges() {
  changesOpen = !changesOpen;
  document.getElementById('changes-list').style.display = changesOpen ? '' : 'none';
  document.getElementById('changes-title').textContent =
    (changesOpen ? '\u25BC ' : '\u25B6 ') + document.getElementById('changes-title').textContent.replace(/^[▼▶] /,'');
}

/* ── Load an RPT ── */
function loadRpt(rpt) {
  currentRpt = rpt;
  currentData = ALL_DATA[rpt];
  selectedJobId = null;

  /* Build adjacency maps */
  nodeMap.clear(); upMap.clear(); downMap.clear();
  currentData.nodes.forEach(function(n){ nodeMap.set(n.id,n); });
  currentData.links.forEach(function(l){
    var s = l.source, t = l.target;
    if (!upMap.has(t)) upMap.set(t,[]);   upMap.get(t).push(s);
    if (!downMap.has(s)) downMap.set(s,[]); downMap.get(s).push(t);
  });

  updateBanner();
  updateChangesList();
  renderGraph();
}

/* ── Banner (savings chips) ── */
function updateBanner() {
  var ss = currentData.sim_stats;
  var chips = document.getElementById('sim-chips');
  var n = currentData.removed_jobs.length;
  if (n === 0) {
    chips.innerHTML = '<span class="sim-chip info">No Consider recs for this RPT</span>';
    return;
  }
  chips.innerHTML =
    '<span class="sim-chip good">Depth: ' + ss.original_max_depth + ' \u2192 ' + ss.new_max_depth + ' (saved ' + ss.depth_saved + ')</span>' +
    '<span class="sim-chip good">Nodes: ' + ss.original_node_count + ' \u2192 ' + ss.new_node_count + ' (saved ' + ss.nodes_saved + ')</span>' +
    '<span class="sim-chip warn">' + ss.rt_saved_minutes.toFixed(1) + ' min saved</span>' +
    '<span class="sim-chip info">' + ss.consider_rec_count + ' recs applied</span>';
}

/* ── Changes list ── */
function updateChangesList() {
  var list = document.getElementById('changes-list');
  list.innerHTML = '';
  var removedJobs = currentData.removed_jobs || [];
  var jobToRec    = currentData.job_to_rec   || {};

  document.getElementById('changes-title').textContent =
    (changesOpen ? '\u25BC' : '\u25B6') + ' Changes (' + removedJobs.length + ' jobs to remove)';

  if (removedJobs.length === 0) {
    list.innerHTML = '<div style="padding:12px 16px;font-size:11px;color:#888;">No jobs marked for removal.</div>';
    return;
  }

  /* Group by recommendation */
  var byRec = {};
  removedJobs.forEach(function(jid){
    var meta = jobToRec[jid] || {};
    var key  = meta.rec_id || '?';
    if (!byRec[key]) byRec[key] = { rec_id: key, category: meta.category||'', jobs:[] };
    byRec[key].jobs.push(jid);
  });

  Object.keys(byRec).sort().forEach(function(key){
    var grp = byRec[key];
    var gDiv = document.createElement('div'); gDiv.className = 'change-group';
    var hdr  = document.createElement('div'); hdr.className  = 'change-group-hdr';
    var cat  = grp.category;
    /* Category-colour dot */
    var catKey = Object.keys(COLORS).find(function(k){ return cat.indexOf(k+'.')===0; }) || 'OTHER';
    hdr.innerHTML = '<span class="dot" style="width:8px;height:8px;border-radius:50%;background:'+COLORS[catKey]+';flex-shrink:0;display:inline-block;"></span>' +
      '<span>['+grp.rec_id+'] &nbsp; ' + (cat.length>50?cat.substring(0,48)+'\u2026':cat) + ' (' + grp.jobs.length + ')</span>';
    gDiv.appendChild(hdr);

    grp.jobs.sort().forEach(function(jid){
      var n = nodeMap.get(jid);
      var itemColor = n ? (COLORS[n.category]||COLORS.OTHER) : '#888';
      var item = document.createElement('div'); item.className = 'change-item'; item.dataset.id = jid;
      item.innerHTML = '<span class="dot" style="background:'+itemColor+'"></span>' +
        '<div><div class="cname">'+jid+'</div>' +
        (n ? '<div class="crec">'+n.category+' | depth '+n.depth+'</div>' : '') +
        '</div>';
      item.onclick = function(){ highlightChangedJob(jid); };
      gDiv.appendChild(item);
    });
    list.appendChild(gDiv);
  });
}

/* Highlight a removed job both in the list and on the graph */
function highlightChangedJob(jid) {
  document.querySelectorAll('.change-item').forEach(function(el){
    el.classList.toggle('highlighted', el.dataset.id === jid);
  });
  selectJob(jid);
}

/* ── Main render ── */
function renderGraph() {
  gRoot.selectAll("*").remove();
  _linkPaths = null; _nodeG = null;

  var data = currentData;
  var removedSet = new Set(data.removed_jobs || []);
  var width  = document.getElementById('graph-container').clientWidth;
  var height = document.getElementById('graph-container').clientHeight;

  /* Filter nodes */
  var allNodes = data.nodes;
  var visibleNodes = allNodes.filter(function(n){
    if (!showUnlinked && n.depth === -1) return false;
    if (mode === 'after' && removedSet.has(n.id)) return false;
    if (difwModeBa === 2 && isDifwNodeBa(n)) return false;
    if (disabledModeBa === 2 && isDisabledNodeBa(n)) return false;
    return true;
  });
  var visibleIds = new Set(visibleNodes.map(function(n){ return n.id; }));
  var visibleLinks = data.links.filter(function(l){
    var s = l.source, t = l.target;
    return visibleIds.has(s) && visibleIds.has(t);
  });

  /* Stats count */
  document.getElementById('stats-count').textContent =
    visibleNodes.length + ' nodes \u00b7 ' + visibleLinks.length + ' edges' +
    (removedSet.size > 0 ? ' \u00b7 ' + removedSet.size + ' will be removed' : '');

  /* Layout by depth column */
  var columns = {};
  visibleNodes.forEach(function(n){
    if (!columns[n.depth]) columns[n.depth] = [];
    columns[n.depth].push(n);
  });
  Object.keys(columns).forEach(function(d){
    columns[d].sort(function(a,b){
      var ca = CAT_ORDER[a.category]!==undefined?CAT_ORDER[a.category]:9;
      var cb = CAT_ORDER[b.category]!==undefined?CAT_ORDER[b.category]:9;
      return ca-cb || a.id.localeCompare(b.id);
    });
  });
  var depths   = Object.keys(columns).map(Number).sort(function(a,b){return a-b;});
  var depthToCol = {}; var colIdx = 0;
  depths.forEach(function(d){ if(d>=0){depthToCol[d]=colIdx++;} });
  if (columns[-1]) { depthToCol[-1]=colIdx++; }
  var numCols = colIdx;

  Object.keys(columns).forEach(function(d){
    var col = columns[d];
    var ci  = depthToCol[parseInt(d)];
    var x   = PAD_LEFT + ci * COL_WIDTH;
    col.forEach(function(n,i){
      n._x = x + NODE_W/2;
      n._y = PAD_TOP + 24 + i * ROW_HEIGHT + NODE_H/2;
    });
  });

  var maxColLen = d3.max(Object.values(columns), function(c){ return c.length; }) || 1;
  var totalW = PAD_LEFT + numCols * COL_WIDTH + 60;
  var totalH = PAD_TOP + 24 + maxColLen * ROW_HEIGHT + 40;

  /* Depth column headers */
  gRoot.selectAll('.depth-col-label').data(depths).join('text')
    .attr('class','depth-col-label')
    .attr('x', function(d){ return PAD_LEFT + depthToCol[d]*COL_WIDTH + NODE_W/2; })
    .attr('y', PAD_TOP)
    .text(function(d){ return (d===-1?'Unlinked':'Depth '+d)+' ('+columns[d].length+')'; });

  /* Arrow marker */
  gRoot.append('defs').append('marker')
    .attr('id','arrow').attr('viewBox','0 -4 8 8')
    .attr('refX',8).attr('refY',0).attr('markerWidth',6).attr('markerHeight',6)
    .attr('orient','auto')
    .append('path').attr('d','M0,-3.5L8,0L0,3.5').attr('fill','#6a7acc');

  /* Links */
  var linkData = visibleLinks.map(function(l){
    return { source: nodeMap.get(l.source), target: nodeMap.get(l.target) };
  }).filter(function(l){ return l.source && l.target; });

  _linkPaths = gRoot.append('g').selectAll('path')
    .data(linkData).join('path')
    .attr('class', function(l){
      var cls = 'link-path';
      if (mode==='diff' && (removedSet.has(l.source.id) || removedSet.has(l.target.id)))
        cls += ' link-removed';
      return cls;
    })
    .attr('marker-end','url(#arrow)')
    .attr('d', function(l){
      var sx=l.source._x-NODE_W/2, sy=l.source._y;
      var tx=l.target._x+NODE_W/2, ty=l.target._y;
      var mx=(sx+tx)/2;
      var P2x=Math.max(mx, tx+NODE_W*0.35);
      return 'M'+sx+','+sy+' C'+mx+','+sy+' '+P2x+','+ty+' '+tx+','+ty;
    });

  /* Nodes */
  _nodeG = gRoot.append('g').selectAll('g')
    .data(visibleNodes).join('g')
    .attr('class','node-group')
    .attr('transform', function(d){ return 'translate('+d._x+','+d._y+')'; });

  var isRem = function(d){ return removedSet.has(d.id); };

  _nodeG.append('rect').attr('class','node-box')
    .attr('width',NODE_W).attr('height',NODE_H)
    .attr('x',-NODE_W/2).attr('y',-NODE_H/2)
    .attr('fill', function(d){
      var c = d3.color(COLORS[d.category]||COLORS.OTHER);
      c.opacity = (mode==='diff'&&isRem(d)) ? 0.10 : 0.18;
      return c;
    })
    .attr('stroke', function(d){
      if (mode==='diff'&&isRem(d)) return '#ff4444';
      if (disabledModeBa===1&&isDisabledNodeBa(d)) return '#6b7280';
      if (difwModeBa===1&&isDifwNodeBa(d)) return '#f59e0b';
      return COLORS[d.category]||COLORS.OTHER;
    })
    .attr('stroke-width', function(d){ return (mode==='diff'&&isRem(d)) ? 2 : (d.depth===0?2.5:1.2); })
    .attr('stroke-dasharray', function(d){
      if (mode==='diff'&&isRem(d)) return '6,3';
      if (disabledModeBa===1&&isDisabledNodeBa(d)) return '5,3';
      return null;
    })
    .attr('opacity', function(d){
      if (mode==='diff'&&isRem(d)) return 0.55;
      if (disabledModeBa===1&&isDisabledNodeBa(d)) return 0.45;
      return 1;
    });

  _nodeG.append('text').attr('class','job-label').attr('text-anchor','middle')
    .attr('dy', function(d){ return d.primary_sql?-10:4; })
    .attr('fill', function(d){ return (mode==='diff'&&isRem(d))?'#ff8888':'#ddd'; })
    .text(function(d){ return d.id.length>40?d.id.substring(0,38)+'\u2026':d.id; });

  _nodeG.filter(function(d){ return d.primary_sql; })
    .append('text').attr('class','sql-label').attr('text-anchor','middle').attr('dy',2)
    .attr('fill', function(d){ return (mode==='diff'&&isRem(d))?'#ff8888':'#69db7c'; })
    .text(function(d){ return d.primary_sql.length>36?d.primary_sql.substring(0,34)+'\u2026':d.primary_sql; });

  /* Depth badge */
  _nodeG.append('circle').attr('r',8)
    .attr('cx',NODE_W/2-6).attr('cy',-NODE_H/2+6)
    .attr('fill','#1a1a3a').attr('stroke','#666').attr('stroke-width',0.8);
  _nodeG.append('text').attr('class','depth-badge').attr('text-anchor','middle')
    .attr('x',NODE_W/2-6).attr('y',-NODE_H/2+10)
    .text(function(d){ return d.depth!=null?d.depth:'?'; });

  /* ✕ badge for removed nodes in diff mode */
  if (mode==='diff') {
    var remNodes = _nodeG.filter(isRem);
    remNodes.append('circle').attr('r',10)
      .attr('cx',-NODE_W/2+12).attr('cy',-NODE_H/2+12)
      .attr('fill','#cc1111').attr('opacity',0.9);
    remNodes.append('text').attr('class','remove-badge')
      .attr('text-anchor','middle').attr('x',-NODE_W/2+12).attr('y',-NODE_H/2+17)
      .text('\u2715');
  }
  /* DIFW badge */
  if (difwModeBa===1) {
    _nodeG.filter(function(d){ return isDifwNodeBa(d); })
      .append('text').attr('class','difw-badge-ba').attr('text-anchor','start')
      .attr('x',-NODE_W/2+4).attr('y',NODE_H/2-4).text('\u29c6 DIFW');
  }
  /* Disabled badge */
  if (disabledModeBa===1) {
    _nodeG.filter(function(d){ return isDisabledNodeBa(d); })
      .append('text').attr('class','disabled-badge-ba').attr('text-anchor','middle')
      .attr('x',0).attr('y',4).text('\u2296 DISABLED');
  }

  /* Tooltip */
  var tooltip = document.getElementById('tooltip');
  _nodeG.on('mouseover', function(e,d){
    var removedMeta = (currentData.job_to_rec||{})[d.id];
    var remNote = removedMeta ? '<br><strong style="color:#ff8888;">\u2715 Will be removed</strong><br>Rec: '+removedMeta.rec_id+'<br>Target: '+removedMeta.target_table : '';
    var difwNote = isDifwNodeBa(d) ? '<br><span style="color:#f59e0b;font-weight:bold;">&#11041; DIFW Load (STG \u2192 DIM/FCT)</span>' : '';
    var disNote  = isDisabledNodeBa(d) ? '<br><span style="color:#9ca3af;font-weight:bold;">&#8856; DISABLED in TIDAL</span>' : '';
    tooltip.style.display='block';
    tooltip.innerHTML = '<strong>'+d.id+'</strong><br>Category: '+d.category+'<br>Depth: '+d.depth+remNote+difwNote+disNote;
  }).on('mousemove', function(e){
    tooltip.style.left=(e.pageX+12)+'px'; tooltip.style.top=(e.pageY-12)+'px';
  }).on('mouseout', function(){ tooltip.style.display='none'; });

  _nodeG.on('click', function(e,d){ e.stopPropagation(); selectJob(d.id); });
  svg.on('click', function(){ clearSelection(); });

  /* Fit */
  setTimeout(function(){
    var sc = Math.min(0.95, Math.min(width/totalW, height/totalH));
    svg.transition().duration(500).call(zoom.transform,
      d3.zoomIdentity.translate(10,10).scale(sc));
  }, 80);
}

/* ── Selection / highlight ── */
function getConnected(jid) {
  var up=new Set(), down=new Set(), q;
  q=[jid]; while(q.length){ var j=q.pop(); (upMap.get(j)||[]).forEach(function(u){ if(!up.has(u)){up.add(u);q.push(u);} }); }
  q=[jid]; while(q.length){ var j=q.pop(); (downMap.get(j)||[]).forEach(function(d){ if(!down.has(d)){down.add(d);q.push(d);} }); }
  return {upstream:up, downstream:down};
}

function selectJob(jid) {
  if (!_nodeG || !_linkPaths) return;
  selectedJobId = jid;
  var conn = getConnected(jid);
  var all  = new Set([jid].concat(Array.from(conn.upstream)).concat(Array.from(conn.downstream)));
  _nodeG.select('.node-box')
    .classed('highlighted', function(d){ return d.id===jid; })
    .attr('opacity', function(d){ return all.has(d.id)?1:0.08; });
  _nodeG.selectAll('text').attr('opacity', function(d){ return all.has(d.id)?1:0.06; });
  _linkPaths.each(function(l){
    var sid=l.source.id, tid=l.target.id;
    var inChain = all.has(sid)&&all.has(tid);
    d3.select(this).classed('highlighted',inChain).classed('dimmed',!inChain);
  });
  showDetail(jid, conn);
}

function clearSelection() {
  selectedJobId = null;
  if (!_nodeG || !_linkPaths) return;
  _nodeG.select('.node-box').classed('highlighted',false).attr('opacity',1);
  _nodeG.selectAll('text').attr('opacity',1);
  _linkPaths.classed('highlighted',false).classed('dimmed',false);
  document.getElementById('detail-panel').style.display='none';
  document.querySelectorAll('.change-item').forEach(function(el){ el.classList.remove('highlighted'); });
}

/* ── Detail panel ── */
function dpField(label, val) {
  return '<div class="dp-field"><div class="dp-label">'+label+'</div><div class="dp-value">'+val+'</div></div>';
}
function showDetail(jid, conn) {
  var n = nodeMap.get(jid); if(!n) return;
  var removedMeta = (currentData.job_to_rec||{})[jid];
  var h = '<h3>'+n.id+'</h3>';

  if (removedMeta) {
    h += '<div class="removed-notice">\u2715 This job will be <strong>removed</strong><br>' +
         'Recommendation: <strong>'+removedMeta.rec_id+'</strong><br>' +
         'Category: '+removedMeta.category+'<br>' +
         'Target table: '+removedMeta.target_table+'</div>';
  }

  h += dpField('Category', n.category);
  h += dpField('Depth', n.depth);
  h += dpField('Status', n.status||'\u2014');
  h += dpField('Lineage Source', n.source||'\u2014');
  if (n.avg_runtime != null) {
    var s=n.avg_runtime;
    h += dpField('Avg Runtime', s>=3600?(s/3600).toFixed(1)+'h':s>=60?(s/60).toFixed(1)+'min':s.toFixed(1)+'s');
  }
  if (n.src_tables&&n.src_tables.length) {
    h += '<div class="dp-field"><div class="dp-label">Source Tables ('+n.src_tables.length+')</div>';
    n.src_tables.forEach(function(t){ h+='<div style="padding:2px 8px;margin:1px 0;background:#1a2a2a;border-left:3px solid #20c997;border-radius:3px;font-size:11px;">'+t+'</div>'; });
    h += '</div>';
  }
  if (n.tgt_tables&&n.tgt_tables.length) {
    h += '<div class="dp-field"><div class="dp-label">Target Tables ('+n.tgt_tables.length+')</div>';
    n.tgt_tables.forEach(function(t){ h+='<div style="padding:2px 8px;margin:1px 0;background:#2a1a2a;border-left:3px solid #ff6b6b;border-radius:3px;font-size:11px;">'+t+'</div>'; });
    h += '</div>';
  }
  if (n.sql_objects&&n.sql_objects.length) {
    h += '<div class="dp-field"><div class="dp-label">SQL Objects</div>';
    n.sql_objects.forEach(function(o){
      h += '<div style="padding:4px 8px;margin:2px 0;background:#1a2a1a;border-left:3px solid #69db7c;border-radius:3px;font-size:11px;"><strong>'+o.name+'</strong>'+(o.type?'<br><span style="color:#888;font-size:10px;">'+o.type+'</span>':'')+'</div>';
    });
    h += '</div>';
  }

  var ups   = Array.from(conn.upstream).sort();
  var downs = Array.from(conn.downstream).sort();

  h += '<div class="dp-field"><div class="dp-label">Upstream ('+ups.length+')</div>';
  ups.slice(0,20).forEach(function(u){
    var un=nodeMap.get(u); var c=un?(COLORS[un.category]||COLORS.OTHER):'#888';
    var remFlag = (currentData.job_to_rec||{})[u] ? ' \u2715' : '';
    h += '<div style="padding:3px 8px;margin:1px 0;background:#1a1a3a;border-left:3px solid '+c+';border-radius:3px;font-size:11px;cursor:pointer;" onclick="selectJob(\''+u+'\')">'+u+remFlag+'</div>';
  });
  if(ups.length>20)h+='<div style="font-size:10px;color:#888;padding:2px 8px">+' +(ups.length-20)+' more</div>';
  h += '</div>';

  h += '<div class="dp-field"><div class="dp-label">Downstream ('+downs.length+')</div>';
  downs.slice(0,20).forEach(function(d){
    var dn=nodeMap.get(d); var c=dn?(COLORS[dn.category]||COLORS.OTHER):'#888';
    var remFlag = (currentData.job_to_rec||{})[d] ? ' \u2715' : '';
    h += '<div style="padding:3px 8px;margin:1px 0;background:#1a1a3a;border-left:3px solid '+c+';border-radius:3px;font-size:11px;cursor:pointer;" onclick="selectJob(\''+d+'\')">'+d+remFlag+'</div>';
  });
  if(downs.length>20)h+='<div style="font-size:10px;color:#888;padding:2px 8px">+'+(downs.length-20)+' more</div>';
  h += '</div>';

  document.getElementById('detail-content').innerHTML = h;
  document.getElementById('detail-panel').style.display = 'block';
}

/* ── Legend ── */
(function(){
  var leg = document.getElementById('legend');
  var h = '<strong>Legend</strong><br>';
  [['RPT','Report Tables'],['FCT','Fact Tables'],['MV','Mat. Views'],['STG','Staging'],['DIM','Dimension'],['OTHER','Other']].forEach(function(p){
    h += '<div class="leg-item"><span class="leg-dot" style="background:'+COLORS[p[0]]+'"></span>'+p[1]+'</div>';
  });
  h += '<div class="leg-item"><span class="leg-dot leg-removed" style="width:12px;height:12px;border-radius:2px;"></span>Will be removed (Diff)</div>';
  leg.innerHTML = h;
})();

/* ── Unlinked toggle ── */
document.getElementById('toggle-unlinked').addEventListener('change', function(e){
  showUnlinked = e.target.checked; renderGraph();
});

/* ── DIFW and Disabled filters ── */
var difwModeBa = 0, disabledModeBa = 0;
function isDifwNodeBa(n) {
  var p=(n.params||'').toUpperCase();
  if (p.indexOf('PKG_GRP_LOAD_DIFW')!==-1) return true;
  return (n.sql_objects||[]).some(function(o){ return o.package&&o.package.toUpperCase().startsWith('PKG_GRP_LOAD_DIFW'); });
}
function isDisabledNodeBa(n) {
  return !!(n.notes && n.notes.toUpperCase().includes('DISABLED'));
}
function _applyBtnStyle(id, mode, labels, bgs, fgs) {
  var btn=document.getElementById(id);
  btn.textContent=labels[mode]; btn.style.background=bgs[mode];
  btn.style.color=fgs[mode]; btn.style.borderColor=mode?bgs[mode]:'';
}
function cycleDifwModeBa() {
  difwModeBa=(difwModeBa+1)%3;
  _applyBtnStyle('btn-difw-ba', difwModeBa,
    ['\u29c6 DIFW Loads','\u29c6 DIFW: Highlighted','\u29c6 DIFW: Hidden'],
    ['','#f59e0b','#7f1d1d'], ['','#000','#fca5a5']);
  renderGraph();
}
function cycleDisabledModeBa() {
  disabledModeBa=(disabledModeBa+1)%3;
  _applyBtnStyle('btn-disabled-ba', disabledModeBa,
    ['\u2296 Disabled Jobs','\u2296 Disabled: Highlighted','\u2296 Disabled: Hidden'],
    ['','#374151','#450a0a'], ['','#d1d5db','#fca5a5']);
  renderGraph();
}

/* ── Stats modal ── */
function _statsFor(nodes, links) {
  var cats={}, totalRt=0, rtCount=0, maxRt=0, sqlObjs=new Set(), tables=new Set();
  var linked=0, unlinked=0, maxDepth=0;
  nodes.forEach(function(n){
    var cat=n.category||'OTHER'; cats[cat]=(cats[cat]||0)+1;
    if(n.avg_runtime!=null){totalRt+=n.avg_runtime;rtCount++;if(n.avg_runtime>maxRt)maxRt=n.avg_runtime;}
    (n.sql_objects||[]).forEach(function(o){if(o.name)sqlObjs.add(o.name);});
    (n.src_tables||[]).forEach(function(t){tables.add(t);});
    (n.tgt_tables||[]).forEach(function(t){tables.add(t);});
    if(n.depth===-1)unlinked++;else if(n.depth>=0)linked++;
    if(n.depth>maxDepth)maxDepth=n.depth;
  });
  return{nodes:nodes.length,edges:links.length,maxDepth:maxDepth,linked:linked,unlinked:unlinked,
         cats:cats,sqlObjs:sqlObjs.size,tables:tables.size,totalRt:totalRt,rtCount:rtCount,maxRt:maxRt};
}

function _fmtRt(sec){
  if(!sec)return '0s';
  if(sec>=3600)return (sec/3600).toFixed(1)+'h';
  if(sec>=60)return (sec/60).toFixed(1)+'m';
  return sec.toFixed(1)+'s';
}

function showStatsModal(){
  var data=currentData;
  var removedSet=new Set(data.removed_jobs||[]);

  /* Apply active DIFW / Disabled filters to both before and after node sets */
  function _applyFilters(nodes) {
    if (difwModeBa===2)      nodes=nodes.filter(function(n){return !isDifwNodeBa(n);});
    if (disabledModeBa===2)  nodes=nodes.filter(function(n){return !isDisabledNodeBa(n);});
    return nodes;
  }
  function _filteredLinks(nodes, links) {
    var ids=new Set(nodes.map(function(n){return n.id;}));
    return links.filter(function(l){return ids.has(l.source)&&ids.has(l.target);});
  }
  var beforeNodes = _applyFilters(data.nodes.slice());
  var afterNodes  = _applyFilters(data.nodes.filter(function(n){return !removedSet.has(n.id);}));
  var b=_statsFor(beforeNodes, _filteredLinks(beforeNodes, data.links));
  var a=_statsFor(afterNodes,  _filteredLinks(afterNodes,  data.links));

  /* Build filter notice */
  var fLabels=[];
  if (difwModeBa===2)      fLabels.push('DIFW hidden');
  if (difwModeBa===1)      fLabels.push('DIFW highlighted');
  if (disabledModeBa===2)  fLabels.push('Disabled hidden');
  if (disabledModeBa===1)  fLabels.push('Disabled highlighted');

  function delta(bv,av,lowerGood){
    var d=av-bv; if(d===0)return '<span class="sc-delta flat">—</span>';
    var good=(lowerGood?d<0:d>0);
    return '<span class="sc-delta '+(good?'good':'bad')+'">'+(d>0?'+':'')+d+'</span>';
  }
  function deltaF(bv,av,lowerGood,dec){
    var d=av-bv; if(d===0)return '<span class="sc-delta flat">—</span>';
    var good=(lowerGood?d<0:d>0);
    return '<span class="sc-delta '+(good?'good':'bad')+'">'+(d>0?'+':'')+d.toFixed(dec||1)+'</span>';
  }

  var h='';

  /* Summary cards */
  var cards=[
    ['Total Nodes', b.nodes, a.nodes, true],
    ['Max Depth',   b.maxDepth, a.maxDepth, true],
    ['Total Edges', b.edges, a.edges, true],
    ['Linked Jobs', b.linked, a.linked, true],
    ['SQL Objects', b.sqlObjs, a.sqlObjs, false],
    ['Tables/Views',b.tables, a.tables, false],
  ];
  h+='<div class="stat-cards-ba">';
  cards.forEach(function(c){
    var d=c[2]-c[1]; var good=(c[3]?d<0:d>0);
    var cls=d===0?'flat':(good?'good':'bad');
    var sign=d>0?'+':'';
    h+='<div class="stat-card-ba">'+'<div class="sc-val">'+c[2]+'</div>'+'<div class="sc-label">'+c[0]+'</div>'+
       '<div class="sc-delta '+cls+'">'+(d===0?'—':sign+d+' vs before')+'</div></div>';
  });
  h+='</div>';

  /* Metrics table */
  h+='<div class="cmp-section-hdr">Key Metrics</div>';
  h+='<table class="cmp-table"><thead><tr><th>Metric</th><th style="text-align:right">Before</th><th style="text-align:right">After</th><th style="text-align:right">Change</th></tr></thead><tbody>';
  [
    ['Total nodes',           b.nodes,    a.nodes,    true,  false],
    ['Max pipeline depth',    b.maxDepth, a.maxDepth, true,  false],
    ['Total edges',           b.edges,    a.edges,    true,  false],
    ['Linked jobs',           b.linked,   a.linked,   true,  false],
    ['Unlinked jobs',         b.unlinked, a.unlinked, true,  false],
    ['SQL objects called',    b.sqlObjs,  a.sqlObjs,  false, false],
    ['Tables / Views',        b.tables,   a.tables,   false, false],
    ['Jobs with runtime data',b.rtCount,  a.rtCount,  false, false],
    ['Total runtime (all jobs)',_fmtRt(b.totalRt),_fmtRt(a.totalRt), true, true],
    ['Max single-job runtime',_fmtRt(b.maxRt),    _fmtRt(a.maxRt),  true, true],
  ].forEach(function(r){
    var chg=r[4]?'<td class="num">—</td>':('<td class="num '+(r[2]-r[1]<0&&r[3]||r[2]-r[1]>0&&!r[3]?'good':'')+'">'+(r[2]-r[1]>0?'+':'')+(r[2]-r[1])+'</td>');
    h+='<tr><td>'+r[0]+'</td><td class="num">'+r[1]+'</td><td class="num">'+r[2]+'</td>'+chg+'</tr>';
  });
  h+='</tbody></table>';

  /* Consider recommendations applied */
  var recs=data.consider_recs||[];
  if(recs.length>0){
    h+='<div class="cmp-section-hdr">Consider Recommendations Applied ('+recs.length+')</div>';
    h+='<table class="cmp-table"><thead><tr><th>ID</th><th>Category</th><th>Target Table</th>'+
       '<th style="text-align:right">Hop Savings</th><th style="text-align:right">RT Saved (min)</th></tr></thead><tbody>';
    recs.forEach(function(r){
      var cat=r.category||''; var catShort=cat.length>45?cat.substring(0,43)+'…':cat;
      h+='<tr><td>'+r.id+'</td><td>'+catShort+'</td><td>'+r.target_table+'</td>'+
         '<td class="num">'+r.verified_hop_savings+'</td>'+
         '<td class="good">'+r.est_runtime_saved_minutes.toFixed(1)+'</td></tr>';
    });
    h+='</tbody></table>';
  }

  document.getElementById('stats-modal-title').textContent=currentRpt+' \u2014 Before vs After Statistics';
  var filterBanner = fLabels.length
    ? '<div style="background:#1a1a2a;border:1px solid #3a3a5a;border-radius:6px;padding:6px 12px;margin-bottom:12px;font-size:11px;color:#ffd43b;">'
      + '\u26a0 Active filters: ' + fLabels.join(' \u00b7 ') + ' \u2014 stats reflect visible nodes only</div>'
    : '';
  document.getElementById('stats-modal-body').innerHTML=filterBanner+h;
  document.getElementById('stats-modal').classList.add('show');
}

/* ── Kick off ── */
loadRpt(currentRpt);
""")
    parts.append("</" + "script>\n</body>\n</html>")
    return "".join(parts)


# ── Main ─────────────────────────────────────────────────────────────────────

def main() -> None:
    print("=" * 65)
    print("Before vs After Hop Reduction Visualizer")
    print("=" * 65)
    print(f"Graph JSON  : {GRAPH_JSON}")
    print(f"Master XLSX : {MASTER_XLSX}")
    print(f"Output HTML : {OUTPUT_HTML}\n")

    if not GRAPH_JSON.exists():
        print(f"ERROR: {GRAPH_JSON} not found. Run generate_tidal_graph.py first.")
        return

    if not MASTER_XLSX.exists():
        print(f"ERROR: Master recommendations file not found:\n  {MASTER_XLSX}")
        return

    print(f"Loading graph JSON ({GRAPH_JSON.name}) ...")
    with open(GRAPH_JSON, "r", encoding="utf-8") as fh:
        all_graphs = json.load(fh)
    print(f"  {len(all_graphs)} RPT graph(s) loaded\n")

    print(f"Loading master recommendations ({MASTER_XLSX.name}) ...")
    master_recs = read_master_recommendations(MASTER_XLSX)
    n_consider = sum(1 for r in master_recs if r["validation_status"] == VS_CONSIDER)
    print(f"  {len(master_recs)} total recs, {n_consider} marked '{VS_CONSIDER}'\n")

    print("Building enhanced graph data ...")
    enhanced_data = build_enhanced_data(all_graphs, master_recs)

    if not enhanced_data:
        print("No data to visualize — check that RPT table names in 'Appears in RPTs' match the graph JSON keys.")
        return

    print(f"\nGenerating HTML ({len(enhanced_data)} RPT(s)) ...")
    html = generate_html(enhanced_data)
    OUTPUT_HTML.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT_HTML, "w", encoding="utf-8") as fh:
        fh.write(html)
    print(f"  Output: {OUTPUT_HTML}")

    print("\nDone! Open the HTML file in a browser to explore:")
    print(f"  {OUTPUT_HTML}")


if __name__ == "__main__":
    main()
