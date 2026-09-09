"""
Generate All_Metadata Object List
=================================
Produces an XLSX with two sheets:
  1. "All_Metadata Objects" — every SQL object (file) in All_Metadata/, with its
     object type classified from the actual CREATE DDL (MV / VIEW / GTT / TABLE / PACKAGE / UNKNOWN).
  2. "Developer Touch List" — objects (packages/procedures AND standalone MV/table
     elimination targets) developers will need to modify or drop, extracted from
     the 'Actionable Work Orders' sheet of output/MASTER_Hop_Reduction_Recommendations.xlsx
     (validation-filtered: only '✅ Consider' / '⏳ Validation Pending' recs).
     '[NO-CODE-CHANGE]' bundle rows are exploded (via the 'Master Recommendations'
     sheet) into one row per MV/table so every elimination candidate is listed
     individually instead of hidden inside one aggregate row. Each row also lists
     'Downstream Dependents' — other All_Metadata objects whose SQL body statically
     references that object (found by scanning All_Metadata/*.sql), so consumers
     not mentioned in the recommendation itself are still surfaced.

Usage:
    python generate_all_metadata_object_list.py
"""

import re
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = Path(__file__).resolve().parent
METADATA_DIR = BASE_DIR / "All_Metadata"
MASTER_XLSX = BASE_DIR / "output" / "MASTER_Hop_Reduction_Recommendations.xlsx"
OUTPUT_XLSX = BASE_DIR / "output" / "All_Metadata_Object_List.xlsx"

_CREATE_RE = re.compile(
    r'CREATE\s+(?:OR\s+REPLACE\s+)?(MATERIALIZED\s+VIEW|GLOBAL\s+TEMPORARY\s+TABLE|VIEW|TABLE|PACKAGE\s+BODY|PACKAGE|PROCEDURE)\s+',
    re.IGNORECASE,
)


def classify_object(sql_file: Path) -> str:
    try:
        text = sql_file.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return "UNKNOWN"
    m = _CREATE_RE.search(text)
    if not m:
        return "UNKNOWN"
    kind = m.group(1).upper().split()
    if kind[0] == "MATERIALIZED":
        return "MATERIALIZED_VIEW"
    if kind[0] == "GLOBAL":
        return "GLOBAL_TEMPORARY_TABLE"
    if kind[0] == "PACKAGE" and len(kind) > 1:
        return "PACKAGE_BODY"
    return kind[0]  # VIEW / TABLE / PACKAGE / PROCEDURE


def collect_metadata_objects() -> list[dict]:
    rows = []
    for sql_file in sorted(METADATA_DIR.glob("*.sql")):
        rows.append({
            "Object Name": sql_file.stem,
            "Object Type": classify_object(sql_file),
            "File Name": sql_file.name,
        })
    return rows


def build_dependents_map() -> dict[str, list[str]]:
    """Static reverse-dependency scan of All_Metadata/*.sql: for every object,
    find every OTHER object whose SQL body references its name (word-boundary
    match). This surfaces consumers not mentioned in the recommendation sheet
    (e.g. a downstream MV that joins to an MV up for elimination), so removing/
    renaming an object also flags who else needs updating.
    """
    files = sorted(METADATA_DIR.glob("*.sql"))
    bodies = {f.stem: f.read_text(encoding="utf-8", errors="ignore") for f in files}
    dependents: dict[str, list[str]] = {name: [] for name in bodies}
    for obj_name in bodies:
        pattern = re.compile(r'\b' + re.escape(obj_name) + r'\b', re.IGNORECASE)
        for other_name, text in bodies.items():
            if other_name == obj_name:
                continue
            if pattern.search(text):
                dependents[obj_name].append(other_name)
    return dependents


def load_master_recommendations() -> dict:
    """Load the 'Master Recommendations' sheet keyed by Master ID, so bundled
    IDs referenced from Actionable Work Orders can be resolved back to their
    Target Table / Category / Validation Status."""
    wb = openpyxl.load_workbook(str(MASTER_XLSX), read_only=True, data_only=True)
    ws = wb["Master Recommendations"]
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter)  # title banner row
    headers = list(next(rows_iter))
    col = {h: i for i, h in enumerate(headers)}

    recs = {}
    for r in rows_iter:
        mid = r[col["Master ID"]]
        if not mid:
            continue
        sql_objs = [o.strip() for o in str(r[col["SQL Objects Called"]] or "").split("\n") if o.strip()]
        recs[str(mid).strip()] = {
            "category": r[col["Category"]],
            "target": r[col["Target Table"]],
            "sql_objects_called": sql_objs,
            "affected_rpts": r[col["Appears in RPTs"]],
            "num_rpts": r[col["# RPTs"]],
            "hop_savings": r[col["Verified Hop Savings"]],
            "est_min": r[col["Est. Time Saved (min)"]],
            "risk": r[col["Risk"]],
            "status": r[col["Validation Status"]],
        }
    wb.close()
    return recs


# Only these categories represent a genuine "eliminate/decommission/migrate this
# MV or table" action where Target Table IS the standalone object to touch.
# Everything else (GTT staging, hardcoded value mapping, cursor folds, circular
# update chains, offset/near-duplicate consolidation, naming violations, etc.)
# is an in-place CODE CHANGE inside the bundled package/procedure — Target Table
# there is just the table referenced by that code, not an independent object for
# developers to separately drop/rename, so it must NOT be exploded into its own row.
_ELIMINATION_CATEGORY_PREFIXES = (
    "e. mv shared across rpts",
    "e. mv elimination candidate",
    "e. mv indicator-chain elimination candidate",
    "e. dead mv",
    "m. orchestration edge without data handoff",
    "j. over-fanout mv",
)


def _is_elimination_category(category: str) -> bool:
    cat_lower = category.strip().lower()
    return any(cat_lower.startswith(p) for p in _ELIMINATION_CATEGORY_PREFIXES)


def _elimination_objects_for(mid: str, master_recs: dict, known_objects: set[str]) -> tuple[str, str, list[str]]:
    """Return (category, recipient_table, elim_objects) for a Master ID, or
    (category, '', []) if this Master ID's category is not a genuine
    elimination/decommission/migration recommendation (see
    _ELIMINATION_CATEGORY_PREFIXES) — such IDs should not get a standalone row.

    'Target Table' means different things by category: for "MV Indicator-Chain"
    recs it's the RECIPIENT that logic gets folded into (excluded from the
    elimination list, with NO fallback to it — if none of the chain's MVs are
    real objects, there's nothing valid to list); for every other eligible
    category (MV Shared/Elimination/Orchestration Edge/Dead MV/Over-Fanout)
    'Target Table' IS the MV/table being eliminated, so it's included.

    Every candidate object is cross-checked against 'known_objects' (the real
    file stems present in All_Metadata/) — objects with no corresponding .sql
    file (e.g. a raw RPT table with no standalone script) are dropped, since
    there is no script for a developer to actually open/modify."""
    mrec = master_recs.get(mid) or {}
    target = mrec.get("target") or ""
    category = str(mrec.get("category") or "")
    if not _is_elimination_category(category):
        return category, "", []
    called = [o for o in mrec.get("sql_objects_called", []) if "." not in o]

    if "indicator-chain" in category.lower():
        recipient = target
        elim_objects = [o for o in called if o.upper() != target.upper()]
    else:
        recipient = ""
        elim_objects = list(dict.fromkeys(([target] if target else []) + called))

    elim_objects = [o for o in elim_objects if o.upper() in known_objects]
    return category, recipient, elim_objects


def collect_dev_touch_list(dependents_map: dict[str, list[str]], known_objects: set[str]) -> list[dict]:
    """Build the developer touch list from the 'Actionable Work Orders' sheet
    (validation-filtered: only '✅ Consider' / '⏳ Validation Pending' recs).

    Rows with a real Package Name are kept as one row per package (a script-level
    PR) AND every bundled Master ID's elimination target(s) (MVs/tables named in
    the Master Recommendations sheet, cross-checked against real All_Metadata
    files via 'known_objects') are ALSO exploded into their own standalone rows —
    bundling an MV-elimination rec inside a package's actionable row must not hide
    that MV as its own actionable object. Rows tagged '[NO-CODE-CHANGE]' are pure
    placeholders (no real package) so only the exploded rows are kept for them."""
    if not MASTER_XLSX.exists():
        print(f"WARNING: {MASTER_XLSX} not found — skipping Developer Touch List sheet")
        return []

    master_recs = load_master_recommendations()

    wb = openpyxl.load_workbook(str(MASTER_XLSX), read_only=True, data_only=True)
    ws = wb["Actionable Work Orders"]
    rows_iter = ws.iter_rows(values_only=True)
    next(rows_iter)  # title banner row
    headers = list(next(rows_iter))
    col = {h: i for i, h in enumerate(headers)}
    all_rows = list(rows_iter)

    # First pass: every real package name used anywhere in this sheet. An object
    # that is itself a package/proc script must never be exploded as an
    # elimination target elsewhere (it's a consumer of the MV, not the MV).
    all_package_names = {
        str(r[col["Package Name"]]).strip().upper()
        for r in all_rows
        if r[col["Package Name"]] and str(r[col["Package Name"]]).strip() not in ("", "[NO-CODE-CHANGE]")
    }

    seen: set[str] = set()
    result = []
    for r in all_rows:
        pkg = r[col["Package Name"]]
        if not pkg or str(pkg).strip() == "":
            continue
        pkg = str(pkg).strip()
        # Section-header rows look like "  RPT_XXX   ▸ Rank N  |  Tier N" and have
        # all other columns None — skip them.
        if r[col["Wave"]] is None and r[col["Scope"]] is None:
            continue

        bundled_ids = [i.strip() for i in str(r[col["Bundled Master IDs"]] or "").split("\n") if i.strip()]

        # Explode every bundled Master ID's elimination target(s) into their own
        # row, regardless of whether this Actionable Work Orders row is a real
        # package or the '[NO-CODE-CHANGE]' placeholder.
        for mid in bundled_ids:
            category, recipient, elim_objects = _elimination_objects_for(mid, master_recs, known_objects)
            mrec = master_recs.get(mid) or {}
            # 'SQL Objects Called' can also list a package/proc that already has
            # its own 'PL/SQL Package / Procedure' row elsewhere in the sheet
            # (a consumer of the MV, not the MV itself) — exclude those globally.
            elim_objects = [o for o in elim_objects if o.upper() not in all_package_names]
            for obj in elim_objects:
                key = f"ELIM::{mid}::{obj}"
                if key in seen:
                    continue
                seen.add(key)
                deps = dependents_map.get(obj, [])
                result.append({
                    "Object / Package Name": obj,
                    "Type": "MV/Table Elimination — No Code Change",
                    "Procedures Touched": "(none — DROP MV / remove from Tidal only)",
                    "Wave": r[col["Wave"]],
                    "Scope": r[col["Scope"]],
                    "Affected RPTs": mrec.get("affected_rpts") or r[col["Affected RPTs"]],
                    "# Actionable Recs": 1,
                    "Categories": category,
                    "Recipient Table (folded into)": recipient,
                    "Downstream Dependents (references this object)": ", ".join(deps),
                    "Hop Savings": mrec.get("hop_savings") or "",
                    "Est. Min Saved": mrec.get("est_min") or "",
                    "Max Risk": mrec.get("risk") or "",
                    "Master IDs": mid,
                    "Validation Status": mrec.get("status") or "",
                })

        if pkg == "[NO-CODE-CHANGE]":
            continue

        if pkg in seen:
            continue
        seen.add(pkg)
        targets = sorted({mrec.get("target") for mid in bundled_ids
                          if (mrec := master_recs.get(mid)) and mrec.get("target")})
        statuses = sorted({mrec.get("status") for mid in bundled_ids
                           if (mrec := master_recs.get(mid)) and mrec.get("status")})
        deps = sorted({d for t in targets for d in dependents_map.get(t, [])})
        result.append({
            "Object / Package Name": pkg,
            "Type": "PL/SQL Package / Procedure",
            "Procedures Touched": r[col["Procedures Touched"]],
            "Wave": r[col["Wave"]],
            "Scope": r[col["Scope"]],
            "Affected RPTs": r[col["Affected RPTs"]],
            "# Actionable Recs": r[col["# Actionable Recs"]],
            "Categories": r[col["Categories"]],
            "Recipient Table (folded into)": "",
            "Downstream Dependents (references this object)": ", ".join(deps),
            "Target Table(s) / MV(s) Bundled": ", ".join(targets),
            "Hop Savings": r[col["Hop Savings"]],
            "Est. Min Saved": r[col["Est. Min Saved"]],
            "Max Risk": r[col["Max Risk"]],
            "Master IDs": ", ".join(bundled_ids),
            "Validation Status": ", ".join(statuses),
        })
    wb.close()
    return _merge_duplicate_elimination_rows(result)


_RISK_RANK = {"HIGH": 3, "MEDIUM": 2, "LOW": 1}


def _to_number(val) -> float:
    try:
        return float(val)
    except (TypeError, ValueError):
        return 0.0


def _merge_duplicate_elimination_rows(rows: list[dict]) -> list[dict]:
    """Collapse duplicate 'MV/Table Elimination' rows for the same object (one
    physical MV/table can be named by several Master IDs, e.g. an Indicator-Chain
    rec for one RPT and a Dead-MV rec for another) into a single row per object,
    merging their Master IDs/Categories/Affected RPTs and keeping the strongest
    Hop Savings / Risk. Package rows are already unique per package and passed
    through unchanged."""
    merged: dict[str, dict] = {}
    package_rows: list[dict] = []
    order: list[str] = []

    for row in rows:
        if row.get("Type") != "MV/Table Elimination — No Code Change":
            package_rows.append(row)
            continue
        obj = row["Object / Package Name"]
        if obj not in merged:
            merged[obj] = dict(row)
            merged[obj]["_waves"] = set()
            merged[obj]["_scopes"] = set()
            merged[obj]["_rpts"] = set()
            merged[obj]["_categories"] = set()
            merged[obj]["_recipients"] = set()
            merged[obj]["_mids"] = set()
            merged[obj]["_statuses"] = set()
            order.append(obj)
        m = merged[obj]
        m["_waves"].add(str(row.get("Wave") or "").strip())
        m["_scopes"].add(str(row.get("Scope") or "").strip())
        m["_rpts"].update(p.strip() for p in str(row.get("Affected RPTs") or "").split(",") if p.strip())
        m["_categories"].add(str(row.get("Categories") or "").strip())
        m["_recipients"].add(str(row.get("Recipient Table (folded into)") or "").strip())
        m["_mids"].add(str(row.get("Master IDs") or "").strip())
        m["_statuses"].add(str(row.get("Validation Status") or "").strip())
        m["Hop Savings"] = max(_to_number(m.get("Hop Savings")), _to_number(row.get("Hop Savings")))
        m["Est. Min Saved"] = max(_to_number(m.get("Est. Min Saved")), _to_number(row.get("Est. Min Saved")))
        if _RISK_RANK.get(str(row.get("Max Risk") or "").upper(), 0) > _RISK_RANK.get(str(m.get("Max Risk") or "").upper(), 0):
            m["Max Risk"] = row.get("Max Risk")

    elim_rows = []
    for obj in order:
        m = merged[obj]
        m["Wave"] = "; ".join(sorted(w for w in m.pop("_waves") if w))
        m["Scope"] = "; ".join(sorted(s for s in m.pop("_scopes") if s))
        m["Affected RPTs"] = ", ".join(sorted(m.pop("_rpts")))
        m["Categories"] = "; ".join(sorted(c for c in m.pop("_categories") if c))
        m["Recipient Table (folded into)"] = ", ".join(sorted(r for r in m.pop("_recipients") if r))
        mids = sorted(m.pop("_mids"))
        m["Master IDs"] = ", ".join(mids)
        m["# Actionable Recs"] = len(mids)
        m["Validation Status"] = ", ".join(sorted(s for s in m.pop("_statuses") if s))
        elim_rows.append(m)

    return package_rows + elim_rows


def write_sheet(ws, headers: list[str], rows: list[dict]):
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(headers, 1):
            val = row.get(h, "")
            if isinstance(val, str) and "\n" in val:
                val = ", ".join(p for p in val.split("\n") if p.strip())
            ws.cell(row=r_idx, column=c_idx, value=val if val is not None else "")
    for c_idx, h in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c_idx)].width = min(50, max(14, len(h) + 4))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    metadata_rows = collect_metadata_objects()
    known_objects = {row["Object Name"].upper() for row in metadata_rows}
    dependents_map = build_dependents_map()
    dev_rows = collect_dev_touch_list(dependents_map, known_objects)

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "All_Metadata Objects"
    write_sheet(ws1, ["Object Name", "Object Type", "File Name"], metadata_rows)

    ws2 = wb.create_sheet("Developer Touch List")
    write_sheet(ws2, [
        "Object / Package Name", "Type", "Procedures Touched", "Wave", "Scope",
        "Affected RPTs", "# Actionable Recs", "Categories", "Recipient Table (folded into)",
        "Downstream Dependents (references this object)",
        "Target Table(s) / MV(s) Bundled", "Hop Savings", "Est. Min Saved", "Max Risk",
        "Master IDs", "Validation Status",
    ], dev_rows)

    OUTPUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(OUTPUT_XLSX))
    print(f"All_Metadata Objects: {len(metadata_rows)} rows")
    print(f"Developer Touch List: {len(dev_rows)} rows")
    print(f"Written: {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
